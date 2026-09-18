import asyncio
import csv
import hashlib
import io
import uuid
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.core.errors import ExternalServiceError, NotFoundError, ValidationError
from app.core.metrics import DOCUMENT_JOBS
from app.db.models import (
    DocumentChunk,
    IntegrationJob,
    ProcessingStatus,
    Project,
    SourceDocument,
)
from app.services.audit import record_audit
from app.services.llm import OpenAIService

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}


class DocumentProcessingService:
    def __init__(self, settings: Settings, llm: OpenAIService) -> None:
        self.settings = settings
        self.llm = llm

    async def process(
        self,
        db: AsyncSession,
        *,
        project_id: uuid.UUID,
        user_id: uuid.UUID,
        filename: str,
        content_type: str,
        content: bytes,
        request_id: str | None = None,
    ) -> tuple[SourceDocument, IntegrationJob, bool]:
        await self._owned_project(db, project_id, user_id)
        safe_name = Path(filename).name
        extension = Path(safe_name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValidationError(
                f"Unsupported document type '{extension or 'unknown'}'. "
                f"Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )
        if not content:
            raise ValidationError("Uploaded document is empty")
        if len(content) > self.settings.max_upload_bytes:
            raise ValidationError(
                f"Document exceeds the {self.settings.max_upload_bytes // (1024 * 1024)} MB limit"
            )

        digest = hashlib.sha256(content).hexdigest()
        existing = await db.scalar(
            select(SourceDocument).where(
                SourceDocument.project_id == project_id, SourceDocument.sha256 == digest
            )
        )
        current_embedding_count = 0
        if existing is not None:
            current_embedding_count = int(
                await db.scalar(
                    select(func.count(DocumentChunk.id)).where(
                        DocumentChunk.document_id == existing.id,
                        DocumentChunk.embedding_model
                        == self.settings.openai_embedding_model,
                    )
                )
                or 0
            )
        if (
            existing
            and existing.status == ProcessingStatus.completed
            and current_embedding_count == existing.chunk_count
            and current_embedding_count > 0
        ):
            job = await db.scalar(
                select(IntegrationJob).where(IntegrationJob.document_id == existing.id)
            )
            if job is None:
                raise ExternalServiceError("Document integration record is missing")
            record_audit(
                db,
                actor_id=user_id,
                action="document_processing.deduplicated",
                resource_type="source_document",
                resource_id=str(existing.id),
                outcome="success",
                details={
                    "sha256": digest,
                    "filename": safe_name,
                    "embedding_model": self.settings.openai_embedding_model,
                },
                request_id=request_id,
            )
            await db.commit()
            return existing, job, True

        document = existing or SourceDocument(
            project_id=project_id,
            user_id=user_id,
            filename=safe_name,
            content_type=content_type or "application/octet-stream",
            sha256=digest,
            size_bytes=len(content),
        )
        if existing is None:
            db.add(document)
            await db.flush()
            job = IntegrationJob(
                document_id=document.id,
                project_id=project_id,
                user_id=user_id,
                idempotency_key=f"document:{project_id}:{digest}",
            )
            db.add(job)
        else:
            job = await db.scalar(
                select(IntegrationJob).where(IntegrationJob.document_id == document.id)
            )
            if job is None:
                job = IntegrationJob(
                    document_id=document.id,
                    project_id=project_id,
                    user_id=user_id,
                    idempotency_key=f"document:{project_id}:{digest}",
                )
                db.add(job)

        document.status = ProcessingStatus.processing
        document.failure_reason = None
        job.status = ProcessingStatus.processing
        record_audit(
            db,
            actor_id=user_id,
            action="document_processing.started",
            resource_type="integration_job",
            resource_id=str(job.id),
            outcome="started",
            details={"filename": safe_name, "sha256": digest, "size_bytes": len(content)},
            request_id=request_id,
        )
        await db.commit()

        last_error: Exception | None = None
        for attempt in range(1, self.settings.document_processing_attempts + 1):
            job.attempts = attempt
            try:
                text = await asyncio.to_thread(self._extract_text, extension, content)
                chunks = self._chunk(text)
                if not chunks:
                    raise ValidationError("No readable text was found in the document")
                vectors = await self.llm.embed_documents(chunks)
                await db.execute(
                    delete(DocumentChunk).where(DocumentChunk.document_id == document.id)
                )
                for index, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
                    db.add(
                        DocumentChunk(
                            document_id=document.id,
                            project_id=project_id,
                            user_id=user_id,
                            chunk_index=index,
                            content=chunk,
                            token_count=max(1, len(chunk) // 4),
                            embedding_model=self.settings.openai_embedding_model,
                            embedding=vector,
                            chunk_metadata={"filename": safe_name, "chunk_index": index},
                        )
                    )
                document.status = ProcessingStatus.completed
                document.chunk_count = len(chunks)
                job.status = ProcessingStatus.completed
                job.last_error = None
                record_audit(
                    db,
                    actor_id=user_id,
                    action="document_processing.completed",
                    resource_type="integration_job",
                    resource_id=str(job.id),
                    outcome="success",
                    details={"attempt": attempt, "chunks": len(chunks)},
                    request_id=request_id,
                )
                await db.commit()
                DOCUMENT_JOBS.labels("completed").inc()
                return document, job, False
            except ValidationError as exc:
                last_error = exc
                break
            except Exception as exc:
                last_error = exc
                job.last_error = str(exc)[:2_000]
                record_audit(
                    db,
                    actor_id=user_id,
                    action="document_processing.retry",
                    resource_type="integration_job",
                    resource_id=str(job.id),
                    outcome="failed",
                    details={"attempt": attempt, "error": str(exc)[:500]},
                    request_id=request_id,
                )
                await db.commit()
                if attempt < self.settings.document_processing_attempts:
                    await asyncio.sleep(min(2 ** (attempt - 1), 4))

        reason = str(last_error or "Document processing failed")[:2_000]
        document.status = ProcessingStatus.failed
        document.failure_reason = reason
        job.status = ProcessingStatus.failed
        job.last_error = reason
        record_audit(
            db,
            actor_id=user_id,
            action="document_processing.failed",
            resource_type="integration_job",
            resource_id=str(job.id),
            outcome="failed",
            details={"attempts": job.attempts, "error": reason[:500]},
            request_id=request_id,
        )
        await db.commit()
        DOCUMENT_JOBS.labels("failed").inc()
        if isinstance(last_error, ValidationError):
            raise last_error
        raise ExternalServiceError("Document processing failed after retries") from last_error

    @staticmethod
    async def _owned_project(
        db: AsyncSession, project_id: uuid.UUID, user_id: uuid.UUID
    ) -> Project:
        project = await db.scalar(
            select(Project).where(Project.id == project_id, Project.user_id == user_id)
        )
        if project is None:
            raise NotFoundError("Project not found")
        return project

    @staticmethod
    def _extract_text(extension: str, content: bytes) -> str:
        stream = io.BytesIO(content)
        if extension == ".pdf":
            return "\n".join(page.extract_text() or "" for page in PdfReader(stream).pages)
        if extension == ".docx":
            return "\n".join(paragraph.text for paragraph in WordDocument(stream).paragraphs)
        if extension == ".xlsx":
            workbook = load_workbook(stream, read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"Sheet: {sheet.title}")
                lines.extend(
                    ", ".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)
                )
            return "\n".join(lines)
        decoded = content.decode("utf-8", errors="replace")
        if extension == ".csv":
            return "\n".join(", ".join(row) for row in csv.reader(io.StringIO(decoded)))
        return decoded

    def _chunk(self, text: str) -> list[str]:
        normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not normalized:
            return []
        chunks: list[str] = []
        start = 0
        size = self.settings.chunk_size_chars
        overlap = self.settings.chunk_overlap_chars
        while start < len(normalized):
            end = min(len(normalized), start + size)
            if end < len(normalized):
                boundary = max(
                    normalized.rfind(". ", start, end),
                    normalized.rfind("\n", start, end),
                )
                if boundary > start + size // 2:
                    end = boundary + 1
            chunk = normalized[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(normalized):
                break
            start = max(start + 1, end - overlap)
        return chunks
